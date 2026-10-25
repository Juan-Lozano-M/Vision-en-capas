"""Reemplaza la portada del informe por una portada APA 7 para estudiante."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Informe_Visión_en_Capas.docx"


def add_page_number(paragraph):
    """Inserta un campo PAGE de Word."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(12)


def add_apa_paragraph(doc, text, bold=False, before=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def has_page_break(element):
    return bool(element.xpath('.//w:br[@w:type="page"]'))


def main():
    doc = Document(DOCX)
    body = doc.element.body
    children = list(body)
    break_index = next((i for i, child in enumerate(children) if has_page_break(child)), None)
    if break_index is None:
        raise RuntimeError("No se encontro el salto de pagina de la portada.")

    # Conserva todo el contenido posterior a la portada.
    remaining = [child for child in children[break_index + 1 :] if child.tag != qn("w:sectPr")]
    for child in children:
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section.header.paragraphs[0])

    add_apa_paragraph(doc, "Deteccion de colores, formas y conteo de objetos", bold=True, before=150)
    add_apa_paragraph(doc, "Prototipo de procesamiento digital de imagenes con Python y OpenCV")
    add_apa_paragraph(doc, "[Nombres de los integrantes]", before=12)
    add_apa_paragraph(doc, "[Nombre de la institucion]")
    add_apa_paragraph(doc, "[Programa o curso]")
    add_apa_paragraph(doc, "[Nombre del docente]")
    add_apa_paragraph(doc, "[Fecha de entrega]")
    doc.add_page_break()

    # Reincorpora el cuerpo existente luego del salto de portada.
    for child in remaining:
        body.append(child)

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()

