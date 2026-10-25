"""Añade el enlace del repositorio y capturas reales de la interfaz al informe."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Informe_Visión_en_Capas.docx"
GITHUB = "https://github.com/Juan-Lozano-M/Vision-en-capas"


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4F46E5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(82, 82, 95)


def add_image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(6.35))
    add_caption(doc, caption)


def main():
    doc = Document(DOCX)
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 1"]
    heading.add_run("Evidencia visual y repositorio")

    p = doc.add_paragraph()
    p.add_run("Repositorio del proyecto: ").bold = True
    add_hyperlink(p, GITHUB, GITHUB)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph("Las siguientes capturas fueron tomadas de la aplicación Streamlit ejecutándose en el entorno local del proyecto. La primera muestra la interfaz y el componente de carga; la segunda muestra la vista general del prototipo con su flujo de análisis.")
    p.paragraph_format.space_after = Pt(8)

    add_image(doc, ROOT / "assets" / "screenshots" / "app_real.png", "Figura 6. Captura real de la interfaz Streamlit y del componente de carga.")
    add_image(doc, ROOT / "assets" / "screenshots" / "app_real_wide.png", "Figura 7. Captura real de la página durante la ejecución del prototipo en una vista amplia.")

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()

