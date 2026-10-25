"""Genera el informe académico del prototipo en formato DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Informe_Visión_en_Capas.docx"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_border(cell, color="D9E2E1", size="6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="172B33", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_border(cell)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text, italic=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.italic = italic
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(82, 105, 108)


def add_figure(doc, filename: str, caption: str, width: float = 6.4):
    """Inserta una figura generada por el prototipo y su pie descriptivo."""
    path = ROOT / "assets" / filename
    if path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)


def add_flow_table(doc):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    labels = [
        ("01", "Imagen original", "Carga y visualización"),
        ("02", "Preprocesamiento", "Redimensionar, suavizar y HSV"),
        ("03", "Extracción", "Máscaras, bordes y contornos"),
        ("04", "Resultado", "Etiquetas y conteos"),
        ("05", "Interpretación", "Lectura y limitaciones"),
    ]
    for cell, (number, title, detail) in zip(table.rows[0].cells, labels):
        shade(cell, "173A42")
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(number)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(127, 224, 190)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(title)
        r2.bold = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(255, 255, 255)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(detail)
        r3.font.size = Pt(7.5)
        r3.font.color.rgb = RGBColor(220, 235, 231)
        cell_border(cell, "5B7B7E")
    add_caption(doc, "Figura 1. Flujo general implementado en la aplicación.")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[idx], "1F5963")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), size=8.8)
            if row_index % 2 == 1:
                shade(cells[idx], "F2F7F6")
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(33, 48, 52)
    for name, size, color in [("Title", 24, "173A42"), ("Heading 1", 16, "173A42"), ("Heading 2", 12, "D45E4C")]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Visión en capas  ·  Prototipo CIPAS")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 120, 120)


def build_report():
    doc = Document()
    configure_document(doc)

    # Portada
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Detección de colores, formas y conteo de objetos")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(25)
    run = subtitle.add_run("Prototipo de procesamiento digital de imágenes con Python y OpenCV")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(212, 94, 76)

    cover = doc.add_table(rows=5, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_data = [
        ("Modalidad", "Trabajo académico en equipo CIPAS"),
        ("Línea seleccionada", "Conteo y clasificación de objetos mediante detección de colores y formas"),
        ("Tecnologías", "Python, OpenCV, NumPy, Matplotlib y Streamlit"),
        ("Integrantes", "[Escribir nombres de los integrantes]"),
        ("Institución y curso", "[Completar información institucional]"),
    ]
    for row, (label, value) in zip(cover.rows, cover_data):
        set_cell_text(row.cells[0], label, bold=True, color="FFFFFF", size=9.5)
        shade(row.cells[0], "173A42")
        set_cell_text(row.cells[1], value, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Informe de desarrollo")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(23, 58, 66)
    doc.add_page_break()

    add_heading(doc, "Resumen", 1)
    add_body(doc, "Este informe documenta el diseño y desarrollo de un prototipo en Python para procesar imágenes, detectar objetos de colores definidos, estimar su forma geométrica y calcular cantidades. La aplicación se construyó con OpenCV para el análisis de imágenes, NumPy para el manejo de matrices, Matplotlib para la visualización de conteos y Streamlit para la interfaz. El flujo de trabajo se presenta de forma explícita como: imagen original, preprocesamiento, extracción de información, resultado e interpretación.")
    add_body(doc, "El prototipo está pensado para imágenes con objetos separados, colores relativamente sólidos y contraste suficiente con el fondo. Las pruebas automatizadas ejecutadas verificaron la clasificación de un triángulo, la detección de un círculo rojo y el resumen de dos objetos sintéticos. Además, la interfaz Streamlit respondió correctamente mediante el servidor local. La evaluación con fotografías reales de la aplicación del equipo debe realizarse cuando dichas imágenes estén disponibles.")

    add_heading(doc, "1 Introducción", 1)
    add_body(doc, "El procesamiento digital de imágenes permite transformar una fotografía en información útil para apoyar tareas de inspección, clasificación y conteo. En un contexto académico, un prototipo pequeño facilita comprender cómo una imagen puede convertirse en datos mediante operaciones matemáticas y algoritmos de visión por computadora.")
    add_body(doc, "Este proyecto desarrolla una aplicación didáctica que recibe una imagen, separa regiones por color, identifica contornos, estima formas geométricas y presenta el resultado de manera visual. La propuesta se alinea con las líneas de conteo de elementos, detección de colores, detección de bordes o formas y clasificación de objetos.")

    add_heading(doc, "2 Planteamiento del problema", 1)
    add_body(doc, "Cuando una persona observa una imagen con varios objetos, puede reconocer colores, formas y cantidades de manera inmediata. Para que un programa realice esa tarea es necesario convertir la imagen a representaciones que permitan distinguir el objeto del fondo y medir características del contorno. Sin un flujo organizado, el análisis resulta difícil de explicar, verificar y reutilizar.")
    add_body(doc, "El problema abordado es, por tanto, cómo construir un prototipo sencillo que procese imágenes obtenidas desde una aplicación, identifique objetos por color y forma, cuente los elementos encontrados y muestre de forma comprensible el camino que conduce al resultado.")

    add_heading(doc, "3 Justificación", 1)
    add_body(doc, "La propuesta es pertinente porque integra programación, álgebra de matrices, procesamiento digital de imágenes y comunicación de resultados en una sola experiencia. La interfaz permite observar las transformaciones intermedias, lo cual ayuda a comprender por qué una detección funciona o falla. Además, la separación entre la lógica de procesamiento y la interfaz facilita que el equipo adapte el prototipo a nuevas fotografías o a un contexto aplicado.")

    add_heading(doc, "4 Objetivos", 1)
    add_heading(doc, "4.1 Objetivo general", 2)
    add_body(doc, "Desarrollar un prototipo funcional en Python que procese imágenes, detecte objetos según su color y forma geométrica, y muestre la cantidad de elementos identificados junto con una interpretación de los resultados.")
    add_heading(doc, "4.2 Objetivos específicos", 2)
    add_bullet(doc, "Implementar un flujo de preprocesamiento que redimensione, suavice y transforme la imagen a un espacio de color adecuado para la segmentación.")
    add_bullet(doc, "Extraer contornos y características geométricas para clasificar figuras y contar objetos por color y forma.")
    add_bullet(doc, "Construir una interfaz visual que presente las cinco etapas del análisis, permita cargar diferentes imágenes y descargue el resultado procesado.")

    add_heading(doc, "5 Marco teórico", 1)
    add_heading(doc, "5.1 Imagen digital y canales de color", 2)
    add_body(doc, "Una imagen digital puede representarse como una matriz cuyos elementos contienen valores de intensidad. En una imagen a color, cada píxel tiene canales que describen su composición. OpenCV suele leer imágenes en el orden BGR, mientras que Matplotlib y muchas interfaces trabajan con RGB.")
    add_heading(doc, "5.2 Espacio HSV", 2)
    add_body(doc, "El espacio HSV separa el tono (Hue), la saturación y el valor o brillo. Esta separación suele ser útil cuando el objetivo es crear una máscara de color, porque permite definir rangos de tono y tolerar variaciones moderadas de iluminación.")
    add_heading(doc, "5.3 Ruido y operaciones morfológicas", 2)
    add_body(doc, "El ruido puede producir regiones pequeñas o discontinuas. El suavizado gaussiano reduce variaciones locales antes de analizar la imagen. Las operaciones morfológicas, como la apertura, ayudan a eliminar pequeños puntos aislados y conservar regiones más consistentes.")
    add_heading(doc, "5.4 Contornos y clasificación geométrica", 2)
    add_body(doc, "Un contorno es una curva que delimita una región. A partir de su perímetro y de una aproximación poligonal es posible estimar el número de vértices. En el prototipo, tres vértices se interpretan como triángulo; cuatro vértices se separan entre cuadrado y rectángulo mediante la proporción de ancho y alto; cinco vértices corresponden a un pentágono. Las figuras con muchos vértices o alta circularidad se interpretan como círculos.")

    add_heading(doc, "6 Herramientas utilizadas", 1)
    add_table(doc, ["Herramienta", "Uso en el prototipo"], [
        ("Python 3.13.2", "Lenguaje principal y ejecución del proyecto."),
        ("OpenCV", "Lectura, conversión de color, filtrado, máscaras, contornos y anotación."),
        ("NumPy", "Matrices de píxeles, rangos HSV y operaciones numéricas."),
        ("Matplotlib", "Gráfico de conteos por categoría."),
        ("Streamlit", "Interfaz gráfica, carga de imágenes y descarga del resultado."),
        ("Pytest", "Pruebas automatizadas de funciones principales."),
    ], widths=[1.55, 5.7])

    add_heading(doc, "7 Metodología", 1)
    add_body(doc, "La metodología se organizó como un pipeline reproducible. La imagen de entrada se conserva para mostrarla al usuario; después se genera una versión redimensionada, suavizada y convertida a HSV. Para cada color configurado se crea una máscara binaria. Los contornos externos que superan el área mínima se analizan y se convierten en registros de detección. Finalmente, se anotan sobre la imagen el identificador, el color y la forma, y se construyen los conteos para la interfaz.")
    add_flow_table(doc)

    add_heading(doc, "8 Desarrollo del prototipo", 1)
    add_heading(doc, "8.1 Imagen original", 2)
    add_body(doc, "La primera etapa permite cargar archivos PNG, JPG, JPEG o WEBP desde el panel lateral de Streamlit. Si no se carga una fotografía, la aplicación utiliza una demostración generada localmente para que el flujo pueda explorarse. La imagen se visualiza antes de aplicar transformaciones y se informa la resolución que será analizada.")
    add_figure(doc, "report_original.png", "Figura 2. Imagen original utilizada para documentar el flujo.")
    add_heading(doc, "8.2 Preprocesamiento", 2)
    add_body(doc, "El prototipo limita el ancho máximo de la imagen para conservar una respuesta ágil sin deformar sus proporciones. Luego aplica un desenfoque gaussiano de 5 por 5 píxeles y convierte la imagen BGR a HSV. También calcula una representación en escala de grises para generar una vista de bordes mediante Canny.")
    add_figure(doc, "report_preprocessing.png", "Figura 3. Imagen después del redimensionamiento y suavizado gaussiano.")
    add_heading(doc, "8.3 Extracción de información", 2)
    add_body(doc, "La extracción utiliza rangos HSV para seis colores: rojo, naranja, amarillo, verde, azul y morado. Para cada máscara se ejecuta una operación morfológica de apertura y se localizan contornos externos. Se descartan las regiones cuya área es menor que el umbral definido por el usuario. Para cada región aceptada se calcula área, centroide, color, forma y una confianza aproximada.")
    add_figure(doc, "report_extraction.png", "Figura 4. Bordes detectados y máscara binaria de regiones de color.")
    add_heading(doc, "8.4 Resultado", 2)
    add_body(doc, "La aplicación dibuja el contorno de cada objeto y coloca una etiqueta con su identificador, color y forma. Además, muestra indicadores con el número total de objetos, el número de colores y formas encontrados, una tabla detallada y un gráfico de conteos. La imagen anotada se puede descargar en formato PNG.")
    add_figure(doc, "report_result.png", "Figura 5. Resultado anotado con contornos, etiquetas y conteo.")
    add_heading(doc, "8.5 Interpretación", 2)
    add_body(doc, "La interpretación resume cuántos objetos se encontraron y cómo se distribuyen por color y forma. También explica que el algoritmo combina segmentación HSV, limpieza morfológica y análisis de contornos. Si no se detectan objetos, la interfaz recomienda reducir el área mínima o utilizar una imagen con mayor contraste. Se reportan como limitaciones las sombras, los reflejos, la superposición de objetos y los fondos con colores parecidos.")

    add_heading(doc, "9 Resultados y análisis", 1)
    add_body(doc, "Los resultados que se presentan a continuación corresponden a verificaciones ejecutadas sobre el proyecto y no a una afirmación de desempeño general en fotografías reales.")
    add_table(doc, ["Verificación", "Resultado observado"], [
        ("Clasificación geométrica", "La prueba del contorno triangular identificó la forma como triángulo y produjo una confianza mayor que cero."),
        ("Detección por color", "La prueba sintética de un círculo rojo detectó un único objeto, con color rojo y forma círculo."),
        ("Resumen de conteos", "La prueba con un círculo rojo y un rectángulo verde produjo un total de dos objetos y un conteo de uno para cada color."),
        ("Interfaz", "El servidor Streamlit respondió con HTTP 200 en el puerto local 8501."),
    ], widths=[1.75, 5.5])
    add_body(doc, "Estas verificaciones indican que las funciones principales y la interfaz básica operan en el entorno configurado. No se debe interpretar esta evidencia como una precisión porcentual, ya que no se ejecutó una evaluación estadística con un conjunto de imágenes etiquetadas. Para completar el análisis del trabajo, el equipo debe probar fotografías reales de su aplicación y registrar el total esperado, el total detectado y los casos de error.")

    add_heading(doc, "10 Limitaciones", 1)
    add_bullet(doc, "Los rangos de color son fijos y pueden requerir calibración según la cámara, la iluminación y los materiales fotografiados.")
    add_bullet(doc, "Los objetos superpuestos o parcialmente ocultos pueden formar un solo contorno y producir un conteo menor.")
    add_bullet(doc, "Las sombras, reflejos y fondos con tonalidades similares pueden generar falsos positivos o regiones incompletas.")
    add_bullet(doc, "La clasificación de formas depende de la calidad del contorno y de la aproximación poligonal; no es un modelo de aprendizaje automático.")
    add_bullet(doc, "La confianza mostrada es una medida heurística para fines didácticos y no equivale a una probabilidad estadística validada.")

    add_heading(doc, "11 Conclusiones", 1)
    add_body(doc, "Se desarrolló un prototipo funcional que integra carga de imágenes, preprocesamiento, segmentación por color, detección de contornos, clasificación geométrica, conteo, visualización y descarga del resultado. La interfaz organiza el análisis según el flujo solicitado, lo que facilita explicar el proceso en una exposición académica.")
    add_body(doc, "Las pruebas ejecutadas confirmaron el funcionamiento de las funciones principales con imágenes sintéticas y verificaron que la interfaz Streamlit responde correctamente. El siguiente paso para fortalecer el trabajo consiste en utilizar imágenes reales generadas por la aplicación del equipo, comparar las detecciones con un conteo manual y documentar los aciertos y errores observados.")

    add_heading(doc, "12 Recomendaciones", 1)
    add_bullet(doc, "Tomar fotografías con iluminación uniforme, fondo contrastante y objetos separados entre sí.")
    add_bullet(doc, "Conservar varias imágenes de prueba para comparar cambios en el umbral de área y en la segmentación de color.")
    add_bullet(doc, "Registrar en una tabla el resultado esperado y el resultado del algoritmo para cada imagen real.")
    add_bullet(doc, "Si el proyecto se amplía, incorporar calibración de colores, separación de objetos conectados y métricas de precisión, recall y F1.")

    add_heading(doc, "Referencias", 1)
    refs = [
        "Bradski, G., & Kaehler, A. (2008). Learning OpenCV: Computer vision with the OpenCV library. O'Reilly Media.",
        "OpenCV. (s. f.). Contour features. https://docs.opencv.org/4.x/dd/d49/tutorial_py_contour_features.html",
        "OpenCV. (s. f.). Changing colorspaces. https://docs.opencv.org/4.x/df/d9d/tutorial_py_colorspaces.html",
        "Python Software Foundation. (s. f.). Python 3 documentation. https://docs.python.org/3/",
        "Streamlit. (s. f.). Streamlit documentation. https://docs.streamlit.io/",
        "NumPy developers. (s. f.). NumPy documentation. https://numpy.org/doc/",
        "Matplotlib development team. (s. f.). Matplotlib documentation. https://matplotlib.org/stable/",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(5)
        p.add_run(ref)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()

